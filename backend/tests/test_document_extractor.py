"""
Testes de extração estruturada de blocos por formato — Fase B.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from app.services.document_extractor_service import TextBlock, extract_blocks


# ---------------------------------------------------------------------------
# TXT
# ---------------------------------------------------------------------------

def test_extract_txt_paragraphs(tmp_path: Path) -> None:
    f = tmp_path / "doc.txt"
    f.write_text("Parágrafo A\n\nParágrafo B", encoding="utf-8")
    blocks = extract_blocks(f, "txt")
    assert len(blocks) == 2
    assert all(b.type == "paragraph" for b in blocks)
    assert blocks[0].text == "Parágrafo A"
    assert blocks[1].text == "Parágrafo B"


def test_extract_txt_empty_blocks_skipped(tmp_path: Path) -> None:
    f = tmp_path / "empty.txt"
    f.write_text("   \n\n   \n\n", encoding="utf-8")
    blocks = extract_blocks(f, "txt")
    assert blocks == []


def test_extract_txt_single_paragraph(tmp_path: Path) -> None:
    f = tmp_path / "single.txt"
    f.write_text("Apenas um parágrafo sem quebras duplas.", encoding="utf-8")
    blocks = extract_blocks(f, "txt")
    assert len(blocks) == 1
    assert blocks[0].type == "paragraph"


def test_extract_txt_order_sequential(tmp_path: Path) -> None:
    f = tmp_path / "order.txt"
    f.write_text("Primeiro\n\nSegundo\n\nTerceiro", encoding="utf-8")
    blocks = extract_blocks(f, "txt")
    orders = [b.order for b in blocks]
    assert orders == sorted(orders)


def test_extract_txt_blocks_have_ids(tmp_path: Path) -> None:
    f = tmp_path / "ids.txt"
    f.write_text("A\n\nB\n\nC", encoding="utf-8")
    blocks = extract_blocks(f, "txt")
    ids = [b.id for b in blocks]
    assert len(set(ids)) == 3  # todos únicos


# ---------------------------------------------------------------------------
# HTML
# ---------------------------------------------------------------------------

def test_extract_html_heading_paragraph_list(tmp_path: Path) -> None:
    f = tmp_path / "doc.html"
    f.write_text("<h1>Titulo</h1><p>Corpo</p><ul><li>Item</li></ul>", encoding="utf-8")
    blocks = extract_blocks(f, "html")
    types = [b.type for b in blocks]
    assert "heading" in types
    assert "paragraph" in types
    assert "list_item" in types


def test_extract_html_heading_level(tmp_path: Path) -> None:
    f = tmp_path / "headings.html"
    f.write_text("<h1>H1</h1><h2>H2</h2><h3>H3</h3>", encoding="utf-8")
    blocks = extract_blocks(f, "html")
    levels = [b.metadata.get("level") for b in blocks]
    assert 1 in levels
    assert 2 in levels
    assert 3 in levels


def test_extract_html_blockquote(tmp_path: Path) -> None:
    f = tmp_path / "quote.html"
    f.write_text("<blockquote>Citacao importante</blockquote>", encoding="utf-8")
    blocks = extract_blocks(f, "html")
    assert any(b.type == "quote" for b in blocks)


def test_extract_html_empty_tags_skipped(tmp_path: Path) -> None:
    f = tmp_path / "empty.html"
    f.write_text("<p></p><h1>   </h1><p>Real</p>", encoding="utf-8")
    blocks = extract_blocks(f, "html")
    assert len(blocks) == 1
    assert blocks[0].text == "Real"


# ---------------------------------------------------------------------------
# DOCX (requer python-docx)
# ---------------------------------------------------------------------------

def test_extract_docx_heading_and_paragraph(tmp_path: Path) -> None:
    pytest.importorskip("docx")
    from docx import Document

    doc = Document()
    doc.add_heading("Capítulo 1", level=1)
    doc.add_paragraph("Este é o corpo do documento.")
    out = tmp_path / "doc.docx"
    doc.save(str(out))

    blocks = extract_blocks(out, "docx")
    types = {b.type for b in blocks}
    assert "heading" in types
    assert "paragraph" in types
    heading = next(b for b in blocks if b.type == "heading")
    assert heading.metadata.get("level") == 1


def test_extract_docx_list_item(tmp_path: Path) -> None:
    pytest.importorskip("docx")
    from docx import Document

    doc = Document()
    p = doc.add_paragraph("Item de lista", style="List Paragraph")
    out = tmp_path / "list.docx"
    doc.save(str(out))

    blocks = extract_blocks(out, "docx")
    assert any(b.type == "list_item" for b in blocks)


# ---------------------------------------------------------------------------
# ODT (requer odfpy)
# ---------------------------------------------------------------------------

def test_extract_odt_basic(tmp_path: Path) -> None:
    odf = pytest.importorskip("odf")
    from odf.opendocument import OpenDocumentText
    from odf.text import H, P
    from odf import teletype

    doc = OpenDocumentText()
    h = H(outlinelevel=1)
    h.addText("Seção")
    doc.text.addElement(h)
    p = P()
    p.addText("Parágrafo ODT")
    doc.text.addElement(p)
    out = tmp_path / "doc.odt"
    doc.save(str(out))

    blocks = extract_blocks(out, "odt")
    types = {b.type for b in blocks}
    assert "heading" in types
    assert "paragraph" in types


# ---------------------------------------------------------------------------
# RTF (requer striprtf)
# ---------------------------------------------------------------------------

def test_extract_rtf_basic(tmp_path: Path) -> None:
    pytest.importorskip("striprtf")
    # RTF mínimo válido
    rtf_content = r"{\rtf1\ansi Hello World\par\par Second paragraph}"
    f = tmp_path / "doc.rtf"
    f.write_text(rtf_content, encoding="utf-8")
    blocks = extract_blocks(f, "rtf")
    # Deve extrair pelo menos um bloco de texto
    assert len(blocks) >= 1
    assert all(b.type == "paragraph" for b in blocks)
